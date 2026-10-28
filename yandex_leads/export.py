"""Выгрузка лидов в Excel / CSV / JSON и хранение сырых данных (JSONL)."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from .models import Business, Lead

COLUMNS = [
    ("Приоритет", 10), ("Балл", 7), ("Название", 32), ("Рубрики", 26), ("Город", 14),
    ("Адрес", 40), ("Телефоны", 22), ("Сайт", 30), ("Соцсети", 30), ("Рейтинг", 9),
    ("Отзывов", 9), ("Что предложить", 34), ("Почему", 60), ("Черновик сообщения", 70),
    ("Карточка на Картах", 30), ("Запрос", 18),
]


def lead_row(lead: Lead) -> list:
    b = lead.business
    socials = list(dict.fromkeys(b.social_links + (lead.audit.social_links if lead.audit else [])))
    reviews = b.review_count if b.review_count is not None else b.rating_count
    return [
        lead.priority,
        lead.score,
        b.name,
        ", ".join(b.categories),
        b.city,
        b.address,
        "\n".join(b.phones),
        lead.own_site or "\n".join(b.websites),
        "\n".join(socials),
        b.rating,
        reviews,
        "\n".join(o.service for o in lead.offers),
        "\n".join(f"{o.service}: {'; '.join(o.reasons)}" for o in lead.offers),
        lead.pitch,
        b.maps_url,
        b.query,
    ]


def save_xlsx(leads: list[Lead], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    ws.append([c for c, _ in COLUMNS])
    head_fill = PatternFill("solid", fgColor="1F2937")
    for i, (_, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width

    fills = {"A": "D1FAE5", "B": "FEF3C7", "C": "F3F4F6"}
    wrap = Alignment(vertical="top", wrap_text=True)
    link_cols = {8, 15}
    for lead in leads:
        ws.append(lead_row(lead))
        r = ws.max_row
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).alignment = wrap
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fills[lead.priority])
        for c in link_cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, str) and cell.value.startswith("http") and "\n" not in cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="2563EB", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Сводка")
    summary.append(["Всего лидов", len(leads)])
    summary.append([])
    summary.append(["Приоритет", "Кол-во"])
    for p, n in sorted(Counter(l.priority for l in leads).items()):
        summary.append([p, n])
    summary.append([])
    summary.append(["Услуга", "Кому предложить"])
    for s, n in Counter(o.service for l in leads for o in l.offers).most_common():
        summary.append([s, n])
    summary.append([])
    summary.append(["Рубрика", "Кол-во"])
    for cat, n in Counter(l.business.categories[0] for l in leads if l.business.categories).most_common(30):
        summary.append([cat, n])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 16
    wb.save(path)


def save_csv(leads: list[Lead], path: Path) -> None:
    # utf-8-sig — чтобы Excel сразу открыл кириллицу без «кракозябр».
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([c for c, _ in COLUMNS])
        for lead in leads:
            w.writerow(lead_row(lead))


def save_json(leads: list[Lead], path: Path) -> None:
    data = [
        {
            "priority": l.priority,
            "score": l.score,
            "business": l.business.to_dict(),
            "maps_url": l.business.maps_url,
            "own_site": l.own_site,
            "audit": l.audit.to_dict() if l.audit else None,
            "offers": [{"service": o.service, "weight": o.weight, "reasons": o.reasons} for o in l.offers],
            "pitch": l.pitch,
        }
        for l in leads
    ]
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def save_leads(leads: list[Lead], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        save_xlsx(leads, path)
    elif suffix == ".csv":
        save_csv(leads, path)
    elif suffix == ".json":
        save_json(leads, path)
    else:
        raise SystemExit(f"Неизвестный формат {suffix}: используйте .xlsx, .csv или .json")


def append_raw(path: Path, b: Business) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(b.to_dict(), ensure_ascii=False) + "\n")


def load_raw(path: Path) -> list[Business]:
    """Читает JSONL; более поздние записи по тому же org_id дополняют ранние."""
    from .sources.extract import merge

    found: dict[str, Business] = {}
    with path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            b = Business.from_dict(json.loads(line))
            if b.org_id in found:
                merge(found[b.org_id], b)
            else:
                found[b.org_id] = b
    return list(found.values())
