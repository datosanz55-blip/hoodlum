import asyncio
import json
from pathlib import Path

import httpx
from openpyxl import load_workbook

from yandex_leads.analyzer import ADMIN, AUTO, CONTENT, SEO, SITE, analyze, parse_services
from yandex_leads.cli import main
from yandex_leads.export import append_raw, load_raw, save_leads
from yandex_leads.models import Business, SiteAudit
from yandex_leads.site_audit import analyze_html, audit_site, is_own_site, pick_own_site
from yandex_leads.sources.api import business_from_feature
from yandex_leads.sources.extract import extract_businesses

FIX = Path(__file__).parent / "fixtures"


def load(name):
    return json.loads((FIX / name).read_text(encoding="utf-8"))


def test_extract_businesses_from_maps_json():
    items = {b.org_id: b for b in extract_businesses(load("maps_search.json"), "барбершоп", "Казань")}
    assert set(items) == {"1001", "1002"}  # топоним отброшен
    boroda = items["1001"]
    assert boroda.name == "Барбершоп Борода"
    assert boroda.address == "Россия, Казань, улица Баумана, 10"
    assert boroda.phones == ["+7 (843) 111-22-33", "+7 (999) 000-00-00"]  # слиты дубли
    assert boroda.rating == 4.3 and boroda.review_count == 9
    assert boroda.lat == 55.79 and boroda.lon == 49.12
    assert boroda.city == "Казань" and boroda.query == "барбершоп"
    smile = items["1002"]
    assert smile.websites == ["http://ulybka-kzn.ru"]
    assert smile.social_links == ["https://vk.com/ulybka"]
    assert smile.categories == ["Стоматологическая клиника", "Имплантация зубов"]


def test_api_feature_parsing():
    feats = load("api_response.json")["features"]
    b1, b2 = (business_from_feature(f, "автосервис", "Самара") for f in feats)
    assert b1.org_id == "2001" and b1.websites == ["http://motor63.ru/"]
    assert b1.phones == ["+7 (846) 300-00-00"] and b1.hours == "пн-сб 9:00–20:00"
    assert b2.websites == [] and b2.categories == ["Шиномонтаж"]


def test_own_site_detection():
    assert not is_own_site("https://vk.com/club1")
    assert not is_own_site("https://boroda.taplink.cc")
    assert not is_own_site("n.yclients.com/company/1")
    assert is_own_site("https://www.ulybka-kzn.ru/")
    assert pick_own_site(["https://vk.com/x", "http://site.ru"]) == "http://site.ru"


GOOD_HTML = """<html><head><title>Стоматология Улыбка — Казань</title>
<meta name="viewport" content="width=device-width">
<meta name="description" content="Лечение и имплантация зубов в Казани, приём без очередей">
<script type="application/ld+json">{}</script>
<script src="https://mc.yandex.ru/metrika/tag.js"></script></head>
<body><a href="/blog/">Блог</a><a href="https://wa.me/79990000000">WhatsApp</a>
<script src="//code.jivo.ru/widget/abc"></script>
<a href="https://n.yclients.com/company/1">Записаться онлайн</a>
<a href="https://vk.com/ulybka">VK</a> &copy; 2015-2026</body></html>"""


def test_analyze_html_good_site():
    a = analyze_html(SiteAudit(url="x"), GOOD_HTML)
    assert a.mobile_viewport and a.meta_description and a.structured_data
    assert a.analytics and a.online_booking and a.chat_widget and a.messengers and a.blog
    assert a.title.startswith("Стоматология") and a.copyright_year == 2026
    assert a.social_links == ["https://vk.com/ulybka"]


def test_analyze_html_bad_site():
    a = analyze_html(SiteAudit(url="x"), "<html><head><title></title></head><body>narod.ru © 2016</body></html>")
    assert not any([a.mobile_viewport, a.meta_description, a.analytics, a.online_booking, a.blog])
    assert a.builder == "uCoz" and a.copyright_year == 2016


def test_audit_site_with_mock_transport():
    def handler(request):
        if request.url.host == "dead.ru":
            raise httpx.ConnectError("boom")
        return httpx.Response(200, html=GOOD_HTML)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=True) as c:
            return await audit_site(c, "https://ok.ru.example"), await audit_site(c, "dead.ru")

    ok, dead = asyncio.run(run())
    assert ok.reachable and ok.https and ok.online_booking
    assert not dead.reachable and dead.error == "ConnectError"


def test_rules_no_site_barber():
    b = Business("1", "Борода", categories=["Барбершоп"], phones=["+7"], websites=["https://vk.com/b"],
                 rating=4.3, review_count=9)
    lead = analyze(b, None, me="Команда X")
    services = [o.service for o in lead.offers]
    assert services[0] == SITE
    assert {SITE, SEO, CONTENT, ADMIN, AUTO} <= set(services)
    assert lead.priority == "A"
    assert "Борода" in lead.pitch and lead.pitch.endswith("Команда X")


def test_rules_good_site_gets_low_score():
    b = Business("2", "Улыбка", categories=["Стоматология", "Имплантация"], websites=["http://u.ru"],
                 social_links=["https://vk.com/u", "https://t.me/u"], rating=4.9, review_count=280)
    audit = analyze_html(SiteAudit(url="http://u.ru", reachable=True, https=True), GOOD_HTML)
    lead = analyze(b, audit)
    assert lead.offers == [] and lead.score == 0 and lead.pitch == ""


def test_dead_site_and_services_filter():
    b = Business("3", "Мотор", categories=["Автосервис"], websites=["http://motor.ru"])
    audit = SiteAudit(url="http://motor.ru", reachable=False, error="ConnectError")
    lead = analyze(b, audit, services=parse_services(["seo,auto"]))
    assert {o.service for o in lead.offers} <= {SEO, AUTO}
    assert SITE in {o.service for o in analyze(b, audit).offers}


def test_export_and_cli_analyze(tmp_path):
    raw = tmp_path / "raw.jsonl"
    for b in extract_businesses(load("maps_search.json"), "барбершоп", "Казань"):
        append_raw(raw, b)
    append_raw(raw, Business("1001", "Барбершоп Борода", phones=["+7 (000) 111-11-11"]))
    loaded = {b.org_id: b for b in load_raw(raw)}
    assert "+7 (000) 111-11-11" in loaded["1001"].phones

    out = tmp_path / "leads.xlsx"
    main(["analyze", str(raw), "-o", str(out), "--no-audit", "--me", "Тест"])
    wb = load_workbook(out)
    ws = wb["Лиды"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Приоритет"
    assert rows[1][2] == "Барбершоп Борода"  # самый «горячий» лид сверху
    assert "Сводка" in wb.sheetnames

    for ext in ("csv", "json"):
        p = tmp_path / f"leads.{ext}"
        save_leads([analyze(b) for b in loaded.values()], p)
        assert p.stat().st_size > 0


def test_analyze_html_unquoted_attributes():
    html = ('<html><head><meta name=viewport content="width=device-width">'
            '<meta name=description content="Маникюр и педикюр в центре Казани, запись онлайн">'
            '<title>Студия</title></head></html>')
    a = analyze_html(SiteAudit(url="x"), html)
    assert a.mobile_viewport and a.meta_description


def test_playwright_proxy_parsing():
    from yandex_leads.sources.browser import playwright_proxy

    assert playwright_proxy(None) is None
    assert playwright_proxy("http://u:p%40ss@1.2.3.4:8080") == {
        "server": "http://1.2.3.4:8080", "username": "u", "password": "p@ss",
    }
    assert playwright_proxy("socks5://host:1080") == {"server": "socks5://host:1080"}
    assert playwright_proxy("5.6.7.8:3128") == {"server": "http://5.6.7.8:3128"}


def test_api_search_skips_seen_ids():
    from yandex_leads.sources.api import ApiCollector

    payload = load("api_response.json")
    api = ApiCollector("key")
    api.client = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, json=payload)))
    found = api.search("автосервис", "Самара", limit=10, skip_ids={"2001"})
    assert [b.org_id for b in found] == ["2002"]
